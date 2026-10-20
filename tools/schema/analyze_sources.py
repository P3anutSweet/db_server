"""Profile product identity and order-line grain without altering the supplied Excel files."""
import argparse
from collections import Counter, defaultdict
import datetime as dt
import hashlib
import json
from pathlib import Path
import re
import string

import openpyxl

PRODUCT_LAYOUTS = {
    '雷神全品类-2017前': ('雷神', 3, {'ean': 'B', 'part': 'C', 'description': 'D', 'registered_date': 'E'}),
    '雷神全品类-2017后': ('雷神', 3, {'ean': 'B', 'part': 'C', 'description': 'D', 'effective_date': 'E'}),
    '雷神全品类-2021后': ('雷神', 5, {'ean': 'B', 'part': 'C', 'description': 'D', 'manufacturer': 'E',
                                      'carton_barcode': 'F', 'carton_spec': 'G', 'updated_date': 'H',
                                      'pid': 'J', 'notes': 'K', 'odm': 'L'}),
    '机械师外设': ('机械师', 2, {'ean': 'B', 'part': 'C', 'description': 'D', 'manufacturer': 'E',
                                  'carton_barcode': 'F', 'carton_spec': 'G', 'updated_date': 'H',
                                  'pid': 'J', 'notes': 'K', 'english_description': 'L', 'supplier': 'M'}),
}
EMPTY = {'', '-', '/', 'N/A', 'NA', '无', '暂无', '未申请', '待申请', '未填写', '未编码', '无料号', '保留'}


def serial(v):
    return v.isoformat() if isinstance(v, (dt.datetime, dt.date, dt.time)) else v


def identifier(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        v = str(int(v)) if int(v) == v else str(v)
    text = str(v).strip().translate(str.maketrans(string.ascii_lowercase, string.ascii_uppercase))
    return None if text in EMPTY else text


def ean_valid(v):
    if not v or not re.fullmatch(r'\d{13}', v):
        return False
    return (10 - sum(int(n) * (1 if i % 2 == 0 else 3) for i, n in enumerate(v[:12])) % 10) % 10 == int(v[-1])


def read_products(path):
    book = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    rows, reserved = [], []
    for name, (brand, start, layout) in PRODUCT_LAYOUTS.items():
        sheet = book[name]
        for number in range(start, sheet.max_row + 1):
            row = {key: serial(sheet[f'{col}{number}'].value) for key, col in layout.items()}
            row['ean_formula'] = row['ean'] if sheet[f'{layout["ean"]}{number}'].data_type == 'f' else None
            if row['ean_formula']:
                row['ean'] = cached[name][f'{layout["ean"]}{number}'].value
            row.update(sheet=name, row=number, brand=brand)
            row['part'] = identifier(row['part'])
            row['ean'] = identifier(row['ean'])
            if row['part'] or row['description']:
                rows.append(row)
            elif row['ean'] or row['ean_formula']:
                reserved.append(row)
    pid_sheet = book['产品PID前期汇总']
    pid_rows = [{openpyxl.utils.get_column_letter(c): serial(pid_sheet.cell(r, c).value) for c in range(1, 9)}
                for r in range(4, pid_sheet.max_row + 1) if pid_sheet.cell(r, 4).value]
    book.close()
    cached.close()
    return rows, reserved, pid_rows


def read_orders(path):
    book = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    rows, sheet_info = [], []
    for sheet in book:
        headers = {openpyxl.utils.get_column_letter(c.column): c.value for c in sheet[1] if c.value is not None}
        before = len(rows)
        for number in range(2, sheet.max_row + 1):
            if not any(sheet[f'{col}{number}'].value is not None for col in ['E', 'H', 'L']):
                continue
            if sheet[f'L{number}'].data_type == 'f' and not sheet[f'E{number}'].value and not sheet[f'H{number}'].value:
                continue
            values = {col: serial(sheet[f'{col}{number}'].value) for col in headers}
            cache = {col: serial(cached[sheet.title][f'{col}{number}'].value) for col in ['L', 'M', 'N']}
            rows.append({'sheet': sheet.title, 'row': number, 'part': identifier(values.get('E')),
                         'ean': identifier(values.get('F')), 'brand': values.get('C'), 'category': values.get('D'),
                         'values': values, 'cached_quantities': cache})
        sheet_info.append({'sheet': sheet.title, 'line_count': len(rows) - before, 'headers': headers,
                           'merged_cells': [str(x) for x in sheet.merged_cells.ranges]})
    book.close()
    cached.close()
    return rows, sheet_info


def reference(row):
    return f"{row['sheet']}!{row['row']}"


def analyze(products, reserved, pid_rows, orders, sheets):
    sku, eans, branded = defaultdict(list), defaultdict(list), defaultdict(list)
    for p in products:
        if p['part']:
            sku[p['part']].append(p)
            branded[(p['brand'], p['part'])].append(p)
        if p['ean']:
            eans[p['ean']].append(p)
    issues = []

    def issue(code, row, detail):
        issues.append({'code': code, 'source': reference(row), 'part': row.get('part'), 'detail': detail})

    missing_part = sum(not p['part'] for p in products)
    bad_ean = [p for p in products if p['ean'] and not ean_valid(p['ean'])]
    formula_ean_missing = [p for p in products if p['ean_formula'] and p['ean'] is None]
    for p in formula_ean_missing:
        issue('product_ean_formula_missing_cache', p, 'EAN is a formula without cached value; recalculate a working copy in Excel before importing barcode')
    malformed_parts = [p for p in products if p['part'] and not re.fullmatch(r'[A-Z0-9]{9}', p['part'])]
    for p in malformed_parts:
        issue('nonstandard_product_part', p, 'Part is not nine ASCII letters/digits; preserve raw text and review possible lookalike characters')
    for p in bad_ean:
        issue('invalid_product_ean', p, 'EAN is not valid EAN-13; preserve raw text and quarantine barcode assignment')
    sku_conflicts = {k: v for k, v in branded.items() if len({r['ean'] for r in v if r['ean']}) > 1}
    ean_conflicts = {k: v for k, v in eans.items() if len({(r['brand'], r['part']) for r in v if r['part']}) > 1}
    missing_skus, ambiguous_skus, ean_disagreements, brand_disagreements = [], [], [], []
    missing_qty, missing_qty_cache, status_qty, negative_qty, formulas = [], [], [], [], []
    missing_fields = Counter()
    statuses = Counter()
    parts = Counter()
    source_dates = defaultdict(list)
    for row in orders:
        values, quantities = row['values'], row['cached_quantities']
        statuses[str(values.get('Q'))] += 1
        parts[row['part']] += 1
        for field, col in [('batch', 'B'), ('part', 'E'), ('ean', 'F'), ('order_date', 'I'),
                           ('forecast_submitted', 'J'), ('planned_delivery', 'K'), ('customer_requirements', 'G')]:
            if values.get(col) is None:
                missing_fields[field] += 1
        for col in ['I', 'J', 'K']:
            if values.get(col):
                source_dates[col].append(str(values[col]))
        matches = sku.get(row['part'], [])
        if not matches:
            missing_skus.append(row)
            issue('missing_product_part', row, 'No exact trimmed/uppercased part match in product master')
        else:
            matches_by_brand = [p for p in matches if p['brand'] == row['brand']]
            if not matches_by_brand:
                brand_disagreements.append(row)
                issue('brand_disagreement', row, 'Part exists, but source brand disagrees with product master')
            candidates = matches_by_brand or matches
            if len({(p['brand'], p['part']) for p in candidates}) > 1:
                ambiguous_skus.append(row)
                issue('ambiguous_product_part', row, 'Part matches more than one brand; no automatic link')
            master_eans = {p['ean'] for p in candidates if p['ean']}
            if row['ean'] and row['ean'] not in master_eans:
                ean_disagreements.append(row)
                issue('order_ean_disagreement', row, 'Order EAN differs from all matching master entries')
        if row['ean'] and not ean_valid(row['ean']):
            issue('invalid_order_ean', row, 'Order EAN fails EAN-13 validation')
        for col in ['L', 'M', 'N']:
            if isinstance(values.get(col), str) and values[col].startswith('='):
                formulas.append({'source': reference(row), 'column': col, 'formula': values[col], 'cache': quantities[col]})
                if quantities[col] is None:
                    missing_qty_cache.append(row)
            if isinstance(quantities[col], (int, float)) and quantities[col] < 0:
                negative_qty.append(row)
        if quantities['L'] is None:
            missing_qty.append(row)
        if values.get('Q') == '已完结' and isinstance(quantities['N'], (int, float)) and quantities['N'] > 0:
            status_qty.append(row)
            issue('closed_with_pending_pickup', row, 'Source status is closed but cached pending pickup quantity is positive; no automatic correction')
    # A repeated SKU or source serial number is not an order identifier.
    duplicate_serials = {}
    for sheet in sheets:
        c = Counter(str(o['values'].get('A')) for o in orders if o['sheet'] == sheet['sheet'])
        duplicate_serials[sheet['sheet']] = {k: v for k, v in c.items() if k != 'None' and v > 1}
    return {
        'product': {'observed_rows': len(products), 'rows_by_sheet': dict(Counter(p['sheet'] for p in products)),
                    'reserved_barcode_rows': len(reserved), 'rows_missing_part': missing_part,
                    'distinct_parts': len(sku), 'distinct_brand_parts': len(branded), 'distinct_eans': len(eans),
                    'duplicate_part_groups': sum(len(v) > 1 for v in sku.values()),
                    'brand_part_multiple_ean_groups': len(sku_conflicts), 'ean_multiple_brand_part_groups': len(ean_conflicts),
                    'invalid_ean_rows': len(bad_ean), 'pid_reference_rows': len(pid_rows)},
        'product_formula_review': {'ean_formulas_without_cached_values': len(formula_ean_missing),
                                   'nonstandard_part_rows': len(malformed_parts)},
        'orders': {'line_count': len(orders), 'sheets': sheets, 'distinct_parts': len(parts),
                   'source_status_counts': dict(statuses), 'missing_fields': dict(missing_fields),
                   'unmatched_part_rows': len(missing_skus), 'ambiguous_part_rows': len(ambiguous_skus),
                   'brand_disagreement_rows': len(brand_disagreements), 'ean_disagreement_rows': len(ean_disagreements),
                   'quantity_formula_cells': len(formulas), 'quantity_formula_cells_missing_cache': len(missing_qty_cache),
                   'closed_with_pending_pickup_rows': len(status_qty), 'missing_order_quantity_rows': len(missing_qty),
                   'negative_quantity_rows': len(negative_qty), 'duplicate_serials': duplicate_serials,
                   'quantity_arithmetic_mismatches': sum(
                       all(isinstance(r['cached_quantities'][c], (int, float)) for c in ['L', 'M', 'N'])
                       and r['cached_quantities']['L'] != r['cached_quantities']['M'] + r['cached_quantities']['N']
                       for r in orders),
                   'date_ranges': {k: [min(v), max(v)] for k, v in source_dates.items()}},
        'issues': issues,
        'conflict_examples': {'brand_part_multiple_ean': [{'brand': k[0], 'part': k[1], 'rows': [reference(r) for r in v]} for k,v in list(sku_conflicts.items())[:10]],
                              'ean_multiple_brand_part': [{'ean': k, 'rows': [reference(r) for r in v]} for k,v in list(ean_conflicts.items())[:10]]},
        'quantity_formula_examples': formulas[:12],
    }


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--products', required=True, type=Path)
    p.add_argument('--orders', required=True, type=Path)
    p.add_argument('--private-output', required=True, type=Path)
    args = p.parse_args()
    products, reserved, pid_rows = read_products(args.products)
    orders, sheets = read_orders(args.orders)
    report = analyze(products, reserved, pid_rows, orders, sheets)
    report['sources'] = [{'name': f.name, 'sha256': hashlib.sha256(f.read_bytes()).hexdigest()} for f in [args.products, args.orders]]
    args.private_output.parent.mkdir(parents=True, exist_ok=True)
    args.private_output.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')
    print(json.dumps(report, ensure_ascii=True, default=str))


if __name__ == '__main__':
    main()
