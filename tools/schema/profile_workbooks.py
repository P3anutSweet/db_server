"""Read-only workbook inventory for database design. Output can contain private source samples."""
import argparse
import datetime as dt
import hashlib
import json
from pathlib import Path
import re

import openpyxl
from openpyxl.utils import get_column_letter


def value(v):
    if isinstance(v, str) and re.search(r'密码|口令|password|账号|帐号|登录', v, re.I):
        return '[excluded sensitive note]'
    return v.isoformat() if isinstance(v, (dt.datetime, dt.date, dt.time)) else v


def inventory(path):
    book = openpyxl.load_workbook(path, data_only=False)
    sheets = []
    for sheet in book:
        if sheet.sheet_state != 'visible' or re.search(r'^(psw|password|credentials)$', sheet.title, re.I):
            sheets.append({'name': sheet.title, 'state': sheet.sheet_state, 'excluded': 'non-business or hidden worksheet'})
            continue
        rows = []
        nonempty = 0
        formulas = 0
        for row in sheet:
            filled = [c for c in row if c.value is not None]
            if not filled:
                continue
            nonempty += 1
            formulas += sum(c.data_type == 'f' for c in filled)
            if len(rows) < 18:
                rows.append({'row': row[0].row, 'cells': {get_column_letter(c.column): value(c.value) for c in filled}})
        sheets.append({'name': sheet.title, 'state': sheet.sheet_state, 'rows': sheet.max_row,
                       'columns': sheet.max_column, 'nonempty_rows': nonempty,
                       'formula_cells': formulas, 'merged_cells': [str(x) for x in sheet.merged_cells.ranges],
                       'top_rows': rows})
    book.close()
    return {'filename': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets': sheets}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('files', nargs='+', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    report = [inventory(path) for path in args.files]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')
    for book in report:
        print(json.dumps({'filename': book['filename']}, ensure_ascii=True))
        for sheet in book['sheets']:
            summary = {k:v for k,v in sheet.items() if k not in ['top_rows','merged_cells']}
            summary['header_rows'] = sheet.get('top_rows', [])[:4]
            print(json.dumps(summary, ensure_ascii=True, default=str))


if __name__ == '__main__':
    main()
